"""PV (Solar) and Wind employment modelling: all five MATLAB scripts in one file.

The file is split into parts that run from top to bottom:

    0. Settings        where the input workbook is and where results are saved
    1. Model inputs    constants, cell ranges and scenario lists taken from the MATLAB scripts
    2. Excel helpers   read cell ranges / write sheets (like MATLAB readmatrix, readtable, writecell)
    3. Model formulas  the employment factor matrices A (employment = A * investment)
    4. Part A          employment by JOB ROLE     (Planning.m, Construction.m, OperationAndMaintenance.m)
    5. Part B          employment by PHASE + charts (Solar_plots.m, Wind_plots.m)
    6. Run everything

The calculations copy the MATLAB scripts exactly, including where the scripts
differ from each other (see the notes next to each setting).

Google Colab: paste this whole file into one cell and run it. Colab already
has numpy, pandas, openpyxl and matplotlib installed.
  * USE_GOOGLE_DRIVE = False: you are asked to upload the input workbook, and
    all results are downloaded to your computer as PV_Wind_Power.zip at the end.
  * USE_GOOGLE_DRIVE = True: the input is read from, and the results are
    written to, the Google Drive folders set below.
"""

import os
import re
import shutil
from collections import namedtuple

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.ticker import FuncFormatter
from openpyxl import Workbook, load_workbook
from openpyxl.utils import range_boundaries

try:
    from google.colab import drive, files  # Only available inside Google Colab
    IN_COLAB = True
except ImportError:
    IN_COLAB = False


# =============================================================================
# 0. SETTINGS: edit these to match where your files are
# =============================================================================
USE_GOOGLE_DRIVE = False  # Colab only: True = read/write in Google Drive, False = upload/download

if IN_COLAB and USE_GOOGLE_DRIVE:
    INPUT_FILE = '/content/drive/MyDrive/Coding/Input_data/employment_modelling_input.xlsx'
    OUTPUT_ROOT = '/content/drive/MyDrive/Coding/PV_Wind_Power'
elif IN_COLAB:
    INPUT_FILE = '/content/employment_modelling_input.xlsx'  # Filled in after the upload
    OUTPUT_ROOT = '/content/PV_Wind_Power'
else:
    INPUT_FILE = r'C:\Users\sultan\OneDrive - Majan Council for Foresight Strategic Affairs and Energy\Desktop\Coding\Input_data\employment_modelling_input.xlsx'
    OUTPUT_ROOT = r'C:\Users\sultan\OneDrive - Majan Council for Foresight Strategic Affairs and Energy\Desktop\Coding\PV_Wind_Power'

# Output folders (same layout as the MATLAB scripts)
JOBROLES_DIR = os.path.join(OUTPUT_ROOT, 'employment_per_jobroles', 'Output_data')  # Part A
STEP1_DIR = os.path.join(OUTPUT_ROOT, 'step1')                                      # Part B, steps B1-B2
SCENARIOS_DATA_DIR = os.path.join(OUTPUT_ROOT, 'scenarios_data')                    # Part B, steps B3-B4


# =============================================================================
# 1. MODEL INPUTS
# =============================================================================

# --- Planning & construction curve: sine rise -> flat peak -> cosine decline ---
q = 0.0        # The height of the curve's peak
k = 0.49       # Power of sine in the first part (rate of increase)
l = 1          # Power of cosine in the third part (rate of decrease)
theta = 1.009  # Scaling parameter

# --- Operation & maintenance curve: arctan (rises quickly, then stays flat) ---
steepness_factor = 10
correction_factor = 2800 / 4393

DURATION_CELL = 'I2:I2'  # Project duration d, on the Solar and Wind sheets

# --- Part A: employment by job role -----------------------------------------
# (scenario name, input column); rows 2-38
JOBROLE_SCENARIOS = [
    ('x_s1', 'C'), ('x_s2', 'D'), ('x_s3', 'E'), ('x_s4', 'F'),
    ('x_gov', 'G'),
    ('x_gov_new', 'G'),  # As in the MATLAB scripts, this reads column G again (same as x_gov)
]
JOBROLE_ROWS = (2, 38)
JOBROLE_YEARS_RANGE = 'A2:A38'
JOBROLE_OM_OFFSET = -0.0011  # Only OperationAndMaintenance.m subtracts this

# (output sheet name, curve type, job-role table range per technology)
JOBROLE_PHASES = [
    ('Planning',                  'plc', {'Solar': 'M2:Q25',  'Wind': 'M2:Q13'}),   # Planning.m
    ('Construction',              'plc', {'Solar': 'M26:Q64', 'Wind': 'M14:Q38'}),  # Construction.m
    ('Operation and Maintenance', 'om',  {'Solar': 'M65:Q77', 'Wind': 'M39:Q54'}),  # OperationAndMaintenance.m
]

# --- Part B: employment by phase + charts -----------------------------------
# (scenario name, input column); rows 2 to the technology's last_row
PHASE_SCENARIOS = [
    ('x_s1', 'C'), ('x_s2', 'D'), ('x_s3', 'E'), ('x_s4', 'F'),
    ('x_gov_sm', 'G'),  # Government plan, smoothed
    ('x_gov_sc', 'H'),  # Government plan, scattered
]
PLOT_SCENARIOS = ['x_s1', 'x_s2', 'x_s3', 'x_s4', 'x_gov_sc', 'x_gov_sm']
LEGEND_NAMES = ['Constant Current', 'Raging Storm', 'Rising Tide', 'Shifting Winds', 'Scattered', 'Smoothed']
STAGE_NAMES = ['Project Planning and Development', 'Construction', 'Operation and Maintenance']
STAGE_COLORS = [(74 / 255, 111 / 255, 165 / 255), (209 / 255, 86 / 255, 86 / 255), (214 / 255, 160 / 255, 118 / 255)]


def _rgb(values):
    return [tuple(v / 255 for v in rgb) for rgb in values]


# Where Solar_plots.m and Wind_plots.m differ
TECHNOLOGIES = {
    'Solar': dict(                       # Solar_plots.m
        last_row=37,                     # Investment in C2:C37 ... H2:H37 (36 half-year periods)
        plc_job_range='T2:X3',           # GOV ('T9:X10' for the scenarios)
        om_job_range='T4:X4',            # GOV ('T11:X11' for the scenarios)
        # Solar_plots.m names the planning & construction sheets with hyphens
        plc_sheet_names=['x-s1', 'x-s2', 'x-s3', 'x-s4', 'x-gov-sm', 'x-gov-sc'],
        om_range_for_totals='B7:C37',    # O&M rows used for the scenario charts (B3)
        stretch_to_even_years=True,      # B3: 2025:0.5:2040 made 32 long; data stretched by interpolation
        scenario_colors=_rgb([(85, 168, 104), (244, 123, 32), (72, 133, 237),
                              (219, 68, 55), (171, 71, 188), (66, 133, 244)]),
        title_main='Employment Over Time for Scenarios x-s1 to x-s4',
        title_gov='Employment Over Time for Scenarios x-gov-sc and x-gov-sm',
    ),
    'Wind': dict(                        # Wind_plots.m
        last_row=32,                     # Investment in C2:C32 ... H2:H32 (31 half-year periods)
        plc_job_range='M55:Q56',         # gov ('T2:X3' for the scenarios)
        om_job_range='M57:Q57',          # Gov ('T3:X3' for the scenarios)
        plc_sheet_names=['x_s1', 'x_s2', 'x_s3', 'x_s4', 'x_gov_sm', 'x_gov_sc'],
        om_range_for_totals='B2:C32',
        stretch_to_even_years=False,
        scenario_colors=_rgb([(85, 168, 104), (244, 123, 32), (72, 133, 237),
                              (219, 68, 55), (85, 168, 104), (66, 133, 244)]),
        title_main='Employment Over Time for Scenarios x_s1 to x_s4',
        title_gov='Employment Over Time for Scenarios x_gov_sc and x_gov_sm',
    ),
}
PLC_RANGE_FOR_CHARTS = 'B7:C37'  # Planning (column B) and construction (column C) rows used in B3 and B4
OM_RANGE_FOR_STACKED = 'B2:B32'  # O&M rows used in B4


# =============================================================================
# 2. EXCEL HELPERS
# =============================================================================
Job = namedtuple('Job', 'name psi start length')

_input_cache = {}


def _open_sheet(path, sheet):
    """Open a sheet for reading (the input workbook is only loaded once)."""
    if path == INPUT_FILE:
        if path not in _input_cache:
            _input_cache[path] = load_workbook(path, data_only=True)
        return _input_cache[path][sheet]
    return load_workbook(path, data_only=True)[sheet]


def read_range(path, sheet, cell_range):
    """Raw cell values of a range, e.g. 'M2:Q25', as a list of rows."""
    ws = _open_sheet(path, sheet)
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    return [[ws.cell(r, c).value for c in range(min_col, max_col + 1)]
            for r in range(min_row, max_row + 1)]


def _to_number(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return np.nan  # Empty or text cells become NaN, as in MATLAB
    return float(value)


def read_numbers(path, sheet, cell_range):
    """Like MATLAB readmatrix: a 2-D array of numbers (NaN for empty/text cells)."""
    return np.array([[_to_number(v) for v in row] for row in read_range(path, sheet, cell_range)])


def read_column(path, sheet, cell_range):
    """A one-column range as a 1-D array, e.g. an investment vector 'C2:C38'."""
    return read_numbers(path, sheet, cell_range)[:, 0]


def read_job_table(path, sheet, cell_range):
    """Like MATLAB readtable on a job-role table (columns Var1..Var5).

    Uses Var2 = job name, Var3 = psi (employment factor), Var4 = start (B/M/E),
    Var5 = length (S/M/L). Completely empty rows are skipped, as readtable does.
    """
    jobs = []
    for row in read_range(path, sheet, cell_range):
        if all(v is None or str(v).strip() == '' for v in row):
            continue
        name, psi, start, length = row[1], row[2], row[3], row[4]
        jobs.append(Job(name='' if name is None else str(name),
                        psi=_to_number(psi),
                        start='' if start is None else str(start),
                        length='' if length is None else str(length)))
    return jobs


def _to_excel(value):
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and np.isnan(value):
        return None
    return value


def write_cells(path, sheet, rows):
    """Like MATLAB writecell/writetable with 'Sheet': write rows from cell A1 of
    the sheet, keeping the workbook's other sheets."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    if os.path.exists(path):
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c).value = _to_excel(value)
    wb.save(path)


# =============================================================================
# 3. MODEL FORMULAS
# =============================================================================
def convert_start_length(start, length, d, early_share):
    """Turn a job's start (B/M/E) and length (S/M/L) into a time window.

    early_share is 0.2 in Planning.m / Construction.m and 0.25 in Solar_plots.m / Wind_plots.m.
    """
    start_time = end_time = None
    if start == 'B':    # Beginning
        start_time = 0
        end_time = {'S': early_share * d, 'M': 0.8 * d, 'L': d}.get(length)
    elif start == 'M':  # Middle
        start_time = early_share * d
        end_time = {'S': 0.8 * d, 'M': d}.get(length)
    elif start == 'E':  # End
        start_time = 0.8 * d
        end_time = {'S': d}.get(length)
    if start_time is None or end_time is None:
        # MATLAB stops with "Output argument not assigned" for these combinations
        raise ValueError(f'Unsupported start/length combination: {start!r}/{length!r}')
    return start_time, end_time


def plc_factor_matrix(psi, start_time, end_time, d, num_periods, formula):
    """Employment factor matrix A for planning & construction jobs.

    formula='jobroles' is the version in Planning.m / Construction.m and
    formula='plots' the version in Solar_plots.m / Wind_plots.m. They differ
    only in the peak and decline pieces; with q = 0 and l = 1 both give the same numbers.
    """
    A = np.zeros((num_periods, num_periods))
    for row in range(1, num_periods + 1):
        for col in range(1, row + 1):
            t = row - col + 1  # Periods since the investment was made
            if not (start_time <= t <= end_time):
                continue
            if t <= d / 3:  # Rise
                value = theta * psi * np.sin((3 * np.pi / (2 * d)) * t) ** k
            elif t <= 2 * d / 3:  # Peak
                bump = q * np.sin((3 * np.pi / d) * (t - d / 3))
                if formula == 'jobroles':
                    bump = abs(bump)
                value = bump + theta * psi
            elif t <= d:  # Decline
                cosine = np.cos((3 * np.pi / d) * (t - 2 * d / 3))
                if formula == 'jobroles':
                    value = (psi * (theta - 1) / 2) * cosine ** l + psi * (1 + theta) / 2
                else:
                    value = (psi * (theta - 1) / 2 * cosine) ** l + psi * (1 + theta) / 2
            else:
                continue
            A[row - 1, col - 1] = value
    return A


def om_factor_matrix(psi, start_time, num_periods, offset):
    """Employment factor matrix A for operation & maintenance jobs (arctan curve).

    offset is -0.0011 in OperationAndMaintenance.m and 0 in Solar_plots.m / Wind_plots.m.
    """
    A = np.zeros((num_periods, num_periods))
    for row in range(1, num_periods + 1):
        for col in range(1, row + 1):
            t = row - col + 1 - start_time
            if t >= 0:
                A[row - 1, col - 1] = psi * correction_factor * np.arctan(steepness_factor * t) + offset
    return A


# =============================================================================
# 4. PART A: EMPLOYMENT BY JOB ROLE
#    (Planning.m, Construction.m, OperationAndMaintenance.m)
#    Output: <Tech>_JobRoles_Scenarios.xlsx with sheets Planning, Construction,
#    Operation and Maintenance. In each sheet the six scenarios are stacked
#    vertically, one column per job role.
# =============================================================================
def run_jobrole_phase(tech, phase_name, curve, job_range):
    jobs = read_job_table(INPUT_FILE, tech, job_range)
    first_row, last_row = JOBROLE_ROWS
    scenarios = [(name, read_column(INPUT_FILE, tech, f'{col}{first_row}:{col}{last_row}'))
                 for name, col in JOBROLE_SCENARIOS]

    if curve == 'plc':
        d = read_numbers(INPUT_FILE, tech, DURATION_CELL)[0, 0]
        years = list(read_column(INPUT_FILE, tech, JOBROLE_YEARS_RANGE))  # Years from the input file
    else:
        # O&M: half-year periods labelled 2025, blank, 2026, blank, ...
        num_periods = len(scenarios[0][1])
        years = [2025 + i // 2 if i % 2 == 0 else None for i in range(num_periods)]
    num_periods = len(years)

    all_rows = []
    for scenario_name, x in scenarios:
        # Header row, then one row per period: Scenario | Time Period | Year | job 1 | job 2 | ...
        block = [['Scenario', 'Time Period', 'Year'] + [job.name for job in jobs]]
        block += [[scenario_name, p + 1, years[p]] + [None] * len(jobs) for p in range(num_periods)]

        for j, job in enumerate(jobs):
            if curve == 'plc':
                if job.psi == 0:
                    continue  # Left blank
                start_time, end_time = convert_start_length(job.start, job.length, d, early_share=0.2)
                A = plc_factor_matrix(job.psi, start_time, end_time, d, num_periods, formula='jobroles')
            else:
                A = om_factor_matrix(job.psi, 0, num_periods, offset=JOBROLE_OM_OFFSET)
            y = A @ x
            for p in range(num_periods):
                block[p + 1][3 + j] = y[p]

        all_rows += block

    output_file = os.path.join(JOBROLES_DIR, f'{tech}_JobRoles_Scenarios.xlsx')
    write_cells(output_file, phase_name, all_rows)
    print(f'{tech} - {phase_name}: employment vectors for all scenarios exported to Excel.')


# =============================================================================
# 5. PART B: EMPLOYMENT BY PHASE + CHARTS (Solar_plots.m, Wind_plots.m)
#    B1  Planning & construction totals -> step1/<Tech>_Employment_Output_PLC.xlsx
#    B2  Operation & maintenance totals -> step1/<Tech>_Employment_Output_OM.xlsx
#    B3  Scenario charts with +/-20% bands -> scenarios_data/<Tech>_Scenarios_Plot_Data.xlsx
#    B4  Stacked area chart per scenario  -> scenarios_data/<Tech>_Stacked_Plot_Data.xlsx
# =============================================================================
def step_b1_planning_construction(tech, cfg, x, plc_file):
    jobs = read_job_table(INPUT_FILE, tech, cfg['plc_job_range'])
    d = read_numbers(INPUT_FILE, tech, DURATION_CELL)[0, 0]
    num_periods = len(x['x_s1'])
    print(f'{tech} PLC - Number of Jobs: {len(jobs)}')
    print(f'{tech} PLC - Number of Periods: {num_periods}')

    Y = {name: np.zeros((num_periods, len(jobs))) for name in x}
    A_matrices = [None] * len(jobs)
    for i, job in enumerate(jobs):
        if job.psi == 0:
            continue  # Skip if no employment factor
        start_time, end_time = convert_start_length(job.start, job.length, d, early_share=0.25)
        A = plc_factor_matrix(job.psi, start_time, end_time, d, num_periods, formula='plots')
        A_matrices[i] = A
        for name in Y:
            Y[name][:, i] = A @ x[name]

    # One sheet per scenario: period numbers in column A, one column per job
    header = [None] + [job.name for job in jobs]
    for (name, _), sheet_name in zip(PHASE_SCENARIOS, cfg['plc_sheet_names']):
        rows = [header] + [[p + 1] + list(Y[name][p]) for p in range(num_periods)]
        write_cells(plc_file, sheet_name, rows)

    # Each job's A matrix on its own sheet (sheet names cleaned, max 31 characters)
    for job, A in zip(jobs, A_matrices):
        if A is not None:
            sheet_name = re.sub(r'[:\\/*?\[\]]', '', job.name)[:31]
            write_cells(plc_file, sheet_name, A.tolist())
    print(f'{tech} PLC employment vectors and factor matrices exported.')


def step_b2_operation_maintenance(tech, cfg, x, om_file):
    jobs = read_job_table(INPUT_FILE, tech, cfg['om_job_range'])
    num_periods = len(x['x_gov_sm'])
    print(f'{tech} OM - Number of Jobs: {len(jobs)}')
    print(f'{tech} OM - Number of Periods: {num_periods}')

    Y = {name: np.zeros((num_periods, len(jobs))) for name in x}
    for i, job in enumerate(jobs):
        A = om_factor_matrix(job.psi, 0, num_periods, offset=0.0)  # For OM, start time is zero
        for name in Y:
            Y[name][:, i] = A @ x[name]

    header = [None] + [job.name for job in jobs]
    for name, _ in PHASE_SCENARIOS:
        rows = [header] + [[p + 1] + list(Y[name][p]) for p in range(num_periods)]
        write_cells(om_file, name, rows)
    print(f'{tech} OM employment vectors exported.')


def _fill_missing_with_zero(values):
    return np.where(np.isnan(values), 0.0, values)  # MATLAB fillmissing(..., 'constant', 0)


def _stretch(values, length):
    """MATLAB interp1(linspace(1, length, numel(values)), values, 1:length, 'linear', 'extrap')."""
    if len(values) == length:
        return values
    return np.interp(np.arange(1, length + 1), np.linspace(1, length, len(values)), values)


def step_b3_scenario_charts(tech, cfg, plc_file, om_file):
    years = np.arange(2025, 2040.25, 0.5)  # 2025:0.5:2040 -> 31 values
    if cfg['stretch_to_even_years'] and len(years) % 2 != 0:
        years = np.append(years, years[-1] + 0.5)

    plc_sheet = dict(zip([name for name, _ in PHASE_SCENARIOS], cfg['plc_sheet_names']))
    total_employment = np.zeros((len(years), len(PLOT_SCENARIOS)))
    for i, scenario in enumerate(PLOT_SCENARIOS):
        data_plc = _fill_missing_with_zero(read_numbers(plc_file, plc_sheet[scenario], PLC_RANGE_FOR_CHARTS))
        data_om = _fill_missing_with_zero(read_numbers(om_file, scenario, cfg['om_range_for_totals']))
        planning, construction, op_maint = data_plc[:, 0], data_plc[:, 1], data_om[:, 0]
        if cfg['stretch_to_even_years']:
            planning, construction, op_maint = (_stretch(v, len(years)) for v in (planning, construction, op_maint))
        total_employment[:, i] = planning + construction + op_maint

    if cfg['stretch_to_even_years'] and np.isnan(total_employment).any():
        raise ValueError('total_employment contains NaN values. Check the input data files or ranges.')

    upper_bounds = total_employment * 1.2
    lower_bounds = total_employment * 0.8

    plot_file = os.path.join(SCENARIOS_DATA_DIR, f'{tech}_Scenarios_Plot_Data.xlsx')
    for indices, title, sheet in [([0, 1, 2, 3], cfg['title_main'], 'Scenarios'),
                                  ([4, 5], cfg['title_gov'], 'GOV_Scenarios')]:
        plot_scenario_bands(years, total_employment, upper_bounds, lower_bounds, indices,
                            cfg['scenario_colors'], title)
        # Export plot data: Year, then Max / Med / Min for each scenario
        header = ['Year']
        for i in indices:
            header += [f'Max_{PLOT_SCENARIOS[i]}', f'Med_{PLOT_SCENARIOS[i]}', f'Min_{PLOT_SCENARIOS[i]}']
        rows = [[years[r]] + [v for i in indices
                              for v in (upper_bounds[r, i], total_employment[r, i], lower_bounds[r, i])]
                for r in range(len(years))]
        write_cells(plot_file, sheet, [header] + rows)


def step_b4_stacked_charts(tech, cfg, plc_file, om_file):
    years = 2025 + np.arange(31) * 0.5  # 2025:0.5:2040
    plc_sheet = dict(zip([name for name, _ in PHASE_SCENARIOS], cfg['plc_sheet_names']))
    stack_file = os.path.join(SCENARIOS_DATA_DIR, f'{tech}_Stacked_Plot_Data.xlsx')

    for scenario in PLOT_SCENARIOS:
        data_plc = read_numbers(plc_file, plc_sheet[scenario], PLC_RANGE_FOR_CHARTS)
        data_om = read_numbers(om_file, scenario, OM_RANGE_FOR_STACKED)
        planning, construction, op_maint = data_plc[:, 0], data_plc[:, 1], data_om[:, 0]
        if not (len(planning) == len(construction) == len(op_maint) == len(years)):
            raise ValueError(f'Data length mismatch in scenario {scenario}. Verify data ranges.')

        plot_stacked(years, [planning, construction, op_maint], f'Employment Over Time for Scenario: {scenario}')

        header = ['Year', 'Project_Planning_Development', 'Construction', 'Operation_Maintenance']
        rows = [[years[r], planning[r], construction[r], op_maint[r]] for r in range(len(years))]
        write_cells(stack_file, scenario, [header] + rows)


def run_technology(tech, cfg):
    plc_file = os.path.join(STEP1_DIR, f'{tech}_Employment_Output_PLC.xlsx')
    om_file = os.path.join(STEP1_DIR, f'{tech}_Employment_Output_OM.xlsx')
    # Investment vectors for each scenario
    x = {name: read_column(INPUT_FILE, tech, f'{col}2:{col}{cfg["last_row"]}') for name, col in PHASE_SCENARIOS}

    step_b1_planning_construction(tech, cfg, x, plc_file)
    step_b2_operation_maintenance(tech, cfg, x, om_file)
    step_b3_scenario_charts(tech, cfg, plc_file, om_file)
    step_b4_stacked_charts(tech, cfg, plc_file, om_file)
    print(f'{tech}: plotting complete. All data exported.')


# --- Charts (MATLAB default figure size 560 x 420 px, axes at [0.13 0.3 0.8 0.6]) ---
def _new_figure():
    fig = plt.figure(figsize=(5.6, 4.2))
    ax = fig.add_axes([0.13, 0.3, 0.8, 0.6])
    return fig, ax


def _finish_axes(fig, ax, title, handles):
    ax.set_xlabel('Year')
    ax.set_ylabel('Number of Jobs')
    ax.set_title(title, fontsize=10, fontweight='bold')
    ax.set_xticks(range(2025, 2041, 5))
    ax.set_xlim(2025, 2040)
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f'{v:.0f}'))  # No 10^x exponent
    ax.grid(True)
    # Legend below the axes without a box (2 columns so it fits the figure width)
    fig.legend(handles=handles, loc='upper center', bbox_to_anchor=(0.5, 0.17), ncol=2, frameon=False, fontsize=9)


def plot_scenario_bands(years, total, upper, lower, indices, colors, title):
    fig, ax = _new_figure()
    handles = []
    for i in indices:
        ax.fill_between(years, lower[:, i], upper[:, i], color=colors[i], alpha=0.3, linewidth=0)
        line, = ax.plot(years, total[:, i], color=colors[i], linewidth=1.5, label=LEGEND_NAMES[i])
        ax.plot(years, upper[:, i], '--', color=colors[i], linewidth=1)
        ax.plot(years, lower[:, i], '--', color=colors[i], linewidth=1)
        handles.append(line)
    _finish_axes(fig, ax, title, handles)


def plot_stacked(years, stages, title):
    fig, ax = _new_figure()
    areas = ax.stackplot(years, *stages, colors=STAGE_COLORS, labels=STAGE_NAMES, linewidth=0)
    _finish_axes(fig, ax, title, areas)


# =============================================================================
# 6. RUN EVERYTHING
# =============================================================================
def prepare_input_file():
    """In Colab, mount Google Drive or ask for the input workbook to be uploaded."""
    global INPUT_FILE
    if not IN_COLAB:
        return
    if USE_GOOGLE_DRIVE:
        drive.mount('/content/drive')
    elif not os.path.exists(INPUT_FILE):
        print('Please upload the input workbook (employment_modelling_input.xlsx):')
        uploaded = files.upload()
        INPUT_FILE = os.path.join('/content', next(iter(uploaded)))


def download_results():
    """In Colab (without Google Drive), download all results as one zip file."""
    if IN_COLAB and not USE_GOOGLE_DRIVE:
        zip_path = shutil.make_archive(OUTPUT_ROOT, 'zip', OUTPUT_ROOT)
        files.download(zip_path)


def main():
    plt.close('all')
    prepare_input_file()

    # Part A: employment by job role (Planning, Construction, O&M) for Solar and Wind
    for phase_name, curve, job_ranges in JOBROLE_PHASES:
        for tech in ('Solar', 'Wind'):
            run_jobrole_phase(tech, phase_name, curve, job_ranges[tech])

    # Part B: employment by phase and charts, Solar then Wind
    for tech, cfg in TECHNOLOGIES.items():
        run_technology(tech, cfg)
        plt.show()

    download_results()


if __name__ == '__main__':
    main()
