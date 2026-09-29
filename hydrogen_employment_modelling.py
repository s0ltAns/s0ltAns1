"""Hydrogen sector employment modelling (Python port of the MATLAB script).

The hydrogen sector is modelled as four subsectors: Solar, Wind, Electrolyser
and Ammonia. For each one the script builds employment factor matrices for the
median, upper and lower psi values, then:

    1. Stacked area charts of jobs by subsector for the two government plans
       (smoothed investment x_gov, scattered investment x_gov_cur)
    2. A line chart of total jobs for the four scenarios with upper/lower bounds
    3. Exports the chart data to H2_Sector_New.xlsx (sheets Raw_x_gov,
       Raw_x_gov_cur, Raw_Scenarios) and saves the three charts as .svg files

Google Colab: paste this whole file into one cell and run it. Colab already
has numpy, openpyxl and matplotlib installed.
  * USE_GOOGLE_DRIVE = False: you are asked to upload the input workbook, and
    the results (Excel file + charts) are downloaded as a zip file at the end.
  * USE_GOOGLE_DRIVE = True: the input is read from, and the results are
    written to, the Google Drive folders set below.
"""

import os
import shutil

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.ticker import FuncFormatter
from openpyxl import Workbook, load_workbook
from openpyxl.utils import range_boundaries

try:
    from google.colab import drive, files  # Only available inside Google Colab
    IN_COLAB = True
except ImportError:
    IN_COLAB = False


# =============================================================================
# 0. SETTINGS: edit these to match where your files are
# =============================================================================
USE_GOOGLE_DRIVE = False  # Colab only: True = read/write in Google Drive, False = upload/download

if IN_COLAB and USE_GOOGLE_DRIVE:
    input_file_name = '/content/drive/MyDrive/LIMA_Employment_modelling/Input_data/employment_modelling_input.xlsx'
    output_file_path = '/content/drive/MyDrive/LIMA_Employment_modelling/Output_data/H2_Sector_New.xlsx'
elif IN_COLAB:
    input_file_name = '/content/employment_modelling_input.xlsx'  # Filled in after the upload
    output_file_path = '/content/Output_data/H2_Sector_New.xlsx'
else:
    input_file_name = r'C:\Users\sultan\OneDrive - Majan Council for Foresight Strategic Affairs and Energy\Desktop\LIMA_Employment_modelling\Input_data\employment_modelling_input.xlsx'
    output_file_path = r'C:\Users\sultan\OneDrive\سطح المكتب\LIMA_Employment_modelling\Output_data\H2_Sector_New.xlsx'

# The .svg charts are saved next to the Excel output file
chart_folder = os.path.dirname(output_file_path)


# =============================================================================
# 1. MODEL INPUTS
# =============================================================================
sheet = 'Hydrogen As one Sector'
T = 'B2:B33'            # Time data (sets the number of periods)
year_range = 'A2:A33'   # Year data (read, but not used, as in the MATLAB script)
start_year = 2024.5     # Each period is half a year

# Investment vectors
SCENARIO_RANGES = {
    'x_s1': 'C2:C33',       # Constant Current
    'x_s2': 'D2:D33',       # Raging Storm
    'x_s3': 'E2:E33',       # Rising Tide
    'x_s4': 'F2:F33',       # Shifting Winds
    'x_gov': 'G2:G33',      # Government investment (smoothed)
    'x_gov_cur': 'H2:H33',  # Current government investment (scattered)
}

SUBSECTORS = ['Solar', 'Wind', 'Electrolyser', 'Ammonia']

# psi cells per subsector: median (column U), upper (column V), lower (column T)
# (The GOV psi values are in P2, P3, P5 and P6 - not used, as in the MATLAB script.)
PSI_CELLS = {
    'median': {'Solar': 'U6:U6', 'Wind': 'U7:U7', 'Electrolyser': 'U9:U9', 'Ammonia': 'U10:U10'},
    'upper':  {'Solar': 'V6:V6', 'Wind': 'V7:V7', 'Electrolyser': 'V9:V9', 'Ammonia': 'V10:V10'},
    'lower':  {'Solar': 'T6:T6', 'Wind': 'T7:T7', 'Electrolyser': 'T9:T9', 'Ammonia': 'T10:T10'},
}

# Subsector-specific d and theta (example values, as in the MATLAB script)
d_map = {'Solar': 1, 'Wind': 1, 'Electrolyser': 1, 'Ammonia': 1}
theta_map = {'Solar': 1, 'Wind': 1, 'Electrolyser': 1, 'Ammonia': 1}

# Constants for employment factor generation
q = 0.2   # Peak height of the curve
k = 0.49  # Power of sine in the increasing phase
l = 1     # Power of cosine in the decreasing phase

# Chart colours
AREA_COLORS = [  # Solar, Wind, Electrolyser, Ammonia
    (209 / 255, 86 / 255, 86 / 255),
    (76 / 255, 175 / 255, 80 / 255),
    (214 / 255, 160 / 255, 118 / 255),
    (165 / 255, 105 / 255, 189 / 255),
]
SCENARIOS = [  # (investment vector, name, colour)
    ('x_s1', 'Constant Current', (0, 0.4470, 0.7410)),       # blue
    ('x_s2', 'Raging Storm', (0.8500, 0.3250, 0.0980)),      # orange/red
    ('x_s3', 'Rising Tide', (0.9290, 0.6940, 0.1250)),       # yellowish
    ('x_s4', 'Shifting Winds', (0.4940, 0.1840, 0.5560)),    # purple
]
SCENARIO_EXPORT_NAMES = ['ConstCurrent', 'RagingStorm', 'RisingTide', 'ShiftingWinds']


# =============================================================================
# 2. EXCEL HELPERS
# =============================================================================
_input_cache = {}


def read_matrix(path, sheet_name, cell_range):
    """Like MATLAB readmatrix: numbers in a range (NaN for empty/text cells)."""
    if path not in _input_cache:
        _input_cache[path] = load_workbook(path, data_only=True)
    ws = _input_cache[path][sheet_name]
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    values = [[ws.cell(r, c).value for c in range(min_col, max_col + 1)] for r in range(min_row, max_row + 1)]
    return np.array([[float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else np.nan
                      for v in row] for row in values])


def read_column(path, sheet_name, cell_range):
    return read_matrix(path, sheet_name, cell_range)[:, 0]


def read_value(path, sheet_name, cell_range):
    return read_matrix(path, sheet_name, cell_range)[0, 0]


def write_cells(path, sheet_name, rows):
    """Like MATLAB writecell with 'Sheet': write rows from cell A1, keeping other sheets."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    if os.path.exists(path):
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            if isinstance(value, np.generic):
                value = value.item()
            if isinstance(value, float) and np.isnan(value):
                value = None
            ws.cell(r, c).value = value
    wb.save(path)


# =============================================================================
# 3. MODEL FORMULA
# =============================================================================
def generate_employment_factor(psi_bar, d, num_periods, q, k, l, theta):
    """Employment factor matrix A: sine rise, peak, cosine decline, then psi_bar after d periods."""
    A = np.zeros((num_periods, num_periods))
    for row in range(1, num_periods + 1):
        for col in range(1, num_periods + 1):
            if row >= col:
                t_diff = row - col + 1
                if t_diff < d / 3:
                    # Increasing phase (sine^k)
                    value = theta * psi_bar * abs(np.sin((3 * np.pi / (2 * d)) * t_diff)) ** k
                elif d / 3 <= t_diff < 2 * d / 3:
                    # Middle phase
                    value = q * abs(np.sin((3 * np.pi / d) * (t_diff - d / 3))) + theta * psi_bar
                elif 2 * d / 3 <= t_diff < d:
                    # Decreasing phase (cosine^l)
                    value = (psi_bar * (theta - 1) / 2) * np.cos((3 * np.pi / d) * (t_diff - 2 * d / 3)) ** l \
                        + psi_bar * (1 + theta) / 2
                else:
                    # After d steps, it settles at psi_bar
                    value = psi_bar
                A[row - 1, col - 1] = value
    return A


# =============================================================================
# 4. CHARTS
# =============================================================================
def _format_axes(ax, title, label_size):
    ax.set_title(title, fontsize=11, fontweight='bold')
    ax.set_xlabel('Years', fontsize=label_size)
    ax.set_ylabel('Number of Jobs', fontsize=label_size)
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f'{v:.0f}'))  # No 10^x exponent
    ax.grid(True)


def _legend_below(ax, handles, labels):
    ax.legend(handles, labels, loc='upper center', bbox_to_anchor=(0.5, -0.15), ncol=2,
              frameon=False, fontsize=10)


def plot_stacked_subsectors(years, data, title, svg_name):
    fig, ax = plt.subplots(figsize=(5.6, 4.2))
    areas = ax.stackplot(years, data.T, colors=AREA_COLORS, linewidth=0)
    ax.set_xlim(years[0], years[-1])
    _format_axes(ax, title, label_size=10)
    _legend_below(ax, areas, SUBSECTORS)
    fig.tight_layout()
    fig.savefig(os.path.join(chart_folder, svg_name))


def plot_with_bounds(ax, x, y, upper, lower, color, label):
    ax.fill_between(x, lower, upper, color=color, alpha=0.1, linewidth=0)
    line, = ax.plot(x, y, linewidth=2, color=color, label=label)
    ax.plot(x, upper, '--', linewidth=1, color=color)
    ax.plot(x, lower, '--', linewidth=1, color=color)
    return line


# =============================================================================
# 5. RUN EVERYTHING
# =============================================================================
def prepare_input_file():
    """In Colab, mount Google Drive or ask for the input workbook to be uploaded."""
    global input_file_name
    if not IN_COLAB:
        return
    if USE_GOOGLE_DRIVE:
        drive.mount('/content/drive')
    elif not os.path.exists(input_file_name):
        print('Please upload the input workbook (employment_modelling_input.xlsx):')
        uploaded = files.upload()
        input_file_name = os.path.join('/content', next(iter(uploaded)))


def download_results():
    """In Colab (without Google Drive), download the output folder as a zip file."""
    if IN_COLAB and not USE_GOOGLE_DRIVE:
        zip_path = shutil.make_archive(chart_folder, 'zip', chart_folder)
        files.download(zip_path)


def main():
    plt.close('all')
    prepare_input_file()
    os.makedirs(chart_folder, exist_ok=True)

    # --- Read data ---
    x = {name: read_column(input_file_name, sheet, rng) for name, rng in SCENARIO_RANGES.items()}
    t = read_column(input_file_name, sheet, T)
    years_table = read_column(input_file_name, sheet, year_range)  # Not used further (as in MATLAB)
    psi = {level: {s: read_value(input_file_name, sheet, cell) for s, cell in cells.items()}
           for level, cells in PSI_CELLS.items()}

    # --- Years for plotting ---
    num_periods = len(t)
    years_plot = start_year + np.arange(num_periods) * 0.5  # Each period represents half a year
    full_year_indices = np.where(years_plot >= 2025)[0]      # Only show data from 2025 onward

    # --- Employment factor matrices: A[level][subsector] ---
    A = {level: {s: generate_employment_factor(psi[level][s], d_map[s], num_periods, q, k, l, theta_map[s])
                 for s in SUBSECTORS}
         for level in PSI_CELLS}

    def jobs_by_subsector(level, investment):
        """Columns: Solar, Wind, Electrolyser, Ammonia."""
        return np.column_stack([A[level][s] @ investment for s in SUBSECTORS])

    def total_jobs(level, investment):
        return jobs_by_subsector(level, investment).sum(axis=1)

    # --- 1. Stacked area charts for the two government plans (median psi) ---
    gov_employment_data = jobs_by_subsector('median', x['x_gov'])[full_year_indices, :]
    gov_cur_employment_data = jobs_by_subsector('median', x['x_gov_cur'])[full_year_indices, :]
    plot_stacked_subsectors(years_plot[full_year_indices], gov_employment_data,
                            'Hydrogen sector (smoothed inv.)', 'Employment_Distribution_x_gov.svg')
    plot_stacked_subsectors(years_plot[full_year_indices], gov_cur_employment_data,
                            'Hydrogen sector (scattered inv.)', 'Employment_Distribution_x_gov_cur.svg')

    # --- 2. Total jobs per scenario with upper and lower bounds ---
    totals = {name: {level: total_jobs(level, x[name]) for level in ('median', 'upper', 'lower')}
              for name, _, _ in SCENARIOS}

    fig, ax = plt.subplots(figsize=(5.6, 4.2))
    handles = [plot_with_bounds(ax, years_plot, totals[name]['median'], totals[name]['upper'],
                                totals[name]['lower'], color, label)
               for name, label, color in SCENARIOS]
    _format_axes(ax, 'Hydrogen sector (scenarios)', label_size=12)
    _legend_below(ax, handles, [label for _, label, _ in SCENARIOS])
    ax.set_xlim(2025, 2040)
    fig.tight_layout()
    fig.savefig(os.path.join(chart_folder, 'Scenario_Line_Plot_with_Bounds.svg'))

    # --- 3. Export the chart data to Excel ---
    header = ['Year'] + SUBSECTORS
    for sheet_name, data in [('Raw_x_gov', gov_employment_data), ('Raw_x_gov_cur', gov_cur_employment_data)]:
        rows = [[year] + list(values) for year, values in zip(years_plot[full_year_indices], data)]
        write_cells(output_file_path, sheet_name, [header] + rows)

    # For each scenario: Max (upper bound), Med (median), Min (lower bound)
    header = ['Year']
    for export_name in SCENARIO_EXPORT_NAMES:
        header += [f'{export_name} Max', f'{export_name} Med', f'{export_name} Min']
    rows = []
    for p in range(num_periods):
        row = [years_plot[p]]
        for name, _, _ in SCENARIOS:
            row += [totals[name]['upper'][p], totals[name]['median'][p], totals[name]['lower'][p]]
        rows.append(row)
    write_cells(output_file_path, 'Raw_Scenarios', [header] + rows)
    print(f'Results saved to {chart_folder}')

    plt.show()
    download_results()


if __name__ == '__main__':
    main()
