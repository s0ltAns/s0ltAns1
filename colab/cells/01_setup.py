# =============================================================================
# STEP 1 - SETUP (run this cell first)
# Settings, file upload and the helper functions used by all the other cells.
# =============================================================================
import math
import os
import re
import shutil
from collections import namedtuple

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.ticker import FuncFormatter
from openpyxl import Workbook, load_workbook
from openpyxl.utils import range_boundaries

try:
    from google.colab import drive, files  # Only available inside Google Colab
    IN_COLAB = True
except ImportError:
    IN_COLAB = False

# -----------------------------------------------------------------------------
# Settings
# -----------------------------------------------------------------------------
# False: you upload the input workbook now, and download all results at the end.
# True:  files are read from / written to your Google Drive folder WORK_DIR.
USE_GOOGLE_DRIVE = False

if IN_COLAB and USE_GOOGLE_DRIVE:
    drive.mount('/content/drive')
    WORK_DIR = '/content/drive/MyDrive/LIMA_Employment_modelling'
elif IN_COLAB:
    WORK_DIR = '/content/LIMA_Employment_modelling'
else:
    WORK_DIR = os.path.abspath('LIMA_Employment_modelling')  # When run outside Colab

INPUT_FILE = os.path.join(WORK_DIR, 'Input_data', 'employment_modelling_input.xlsx')
OUTPUT_DIR = os.path.join(WORK_DIR, 'Output_data')
RAW_DATA_FILE = os.path.join(OUTPUT_DIR, 'raw_data.xlsx')  # Every model cell adds its sheet(s) here
os.makedirs(os.path.dirname(INPUT_FILE), exist_ok=True)
os.makedirs(OUTPUT_DIR, exist_ok=True)

# -----------------------------------------------------------------------------
# Upload the input workbook (and, optionally, your previous raw_data.xlsx)
# -----------------------------------------------------------------------------
if IN_COLAB and not USE_GOOGLE_DRIVE:
    print('Select employment_modelling_input.xlsx.\n'
          'Optional: also select your previous raw_data.xlsx - any sheet that no cell below '
          'recreates (e.g. BuildingEnergyEfficiency) is then kept.')
    for name, content in files.upload().items():
        target = RAW_DATA_FILE if 'raw_data' in name.lower() else INPUT_FILE
        with open(target, 'wb') as f:
            f.write(content)
        if os.path.exists(name):
            os.remove(name)  # files.upload() also saves a copy in /content
        print(f'Saved {name} -> {target}')

if not os.path.exists(INPUT_FILE):
    raise FileNotFoundError(f'Input workbook not found: {INPUT_FILE}')

# -----------------------------------------------------------------------------
# Excel helpers (like MATLAB readmatrix / readtable / writecell)
# -----------------------------------------------------------------------------
_workbook_cache = {}


def _open_sheet(path, sheet):
    key = (path, os.path.getmtime(path))
    if key not in _workbook_cache:
        _workbook_cache[key] = load_workbook(path, data_only=True)
    return _workbook_cache[key][sheet]


def read_range(path, sheet, cell_range):
    """Raw cell values of a range such as 'M2:Q25', as a list of rows."""
    ws = _open_sheet(path, sheet)
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    return [[ws.cell(r, c).value for c in range(min_col, max_col + 1)]
            for r in range(min_row, max_row + 1)]


def to_number(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return np.nan  # Empty or text cells become NaN, as in MATLAB
    return float(value)


def read_numbers(path, sheet, cell_range):
    """Like MATLAB readmatrix: a 2-D array of numbers."""
    return np.array([[to_number(v) for v in row] for row in read_range(path, sheet, cell_range)])


def read_column(path, sheet, cell_range):
    return read_numbers(path, sheet, cell_range)[:, 0]


def read_value(path, sheet, cell_range):
    return read_numbers(path, sheet, cell_range)[0, 0]


def make_valid_name(name):
    """Mimic MATLAB's variable naming in readtable ('_s1' -> 'x_s1', 'psi_bar Max' -> 'psi_barMax')."""
    name = str(name).strip()
    name = re.sub(r'\s+(\w)', lambda m: m.group(1).upper(), name)
    name = re.sub(r'\s+', '', name)
    name = re.sub(r'\W', '_', name, flags=re.ASCII)
    if not name or not name[0].isalpha():
        name = 'x' + name
    return name


def read_table(path, sheet, cell_range):
    """Like MATLAB readtable: first row = column names, numeric data below (empty rows skipped)."""
    rows = read_range(path, sheet, cell_range)
    header = [make_valid_name(h) for h in rows[0]]
    data = [[to_number(v) for v in row] for row in rows[1:]
            if any(v is not None and str(v).strip() != '' for v in row)]
    return pd.DataFrame(data, columns=header)


Job = namedtuple('Job', 'name psi start length')


def read_job_table(path, sheet, cell_range):
    """Job-role table without header: Var2 = name, Var3 = psi, Var4 = start (B/M/E), Var5 = length (S/M/L)."""
    jobs = []
    for row in read_range(path, sheet, cell_range):
        if all(v is None or str(v).strip() == '' for v in row):
            continue
        name, psi, start, length = row[1], row[2], row[3], row[4]
        jobs.append(Job('' if name is None else str(name), to_number(psi),
                        '' if start is None else str(start), '' if length is None else str(length)))
    return jobs


def _to_excel(value):
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and np.isnan(value):
        return None
    return value


def write_cells(path, sheet, rows, replace_sheet=False):
    """Like MATLAB writecell with 'Sheet': write rows from cell A1, keeping the other sheets.
    replace_sheet=True clears the sheet first (used for raw_data.xlsx)."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    if os.path.exists(path):
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    if sheet in wb.sheetnames and replace_sheet:
        index = wb.sheetnames.index(sheet)
        wb.remove(wb[sheet])
        ws = wb.create_sheet(sheet, index)
    elif sheet in wb.sheetnames:
        ws = wb[sheet]
    else:
        ws = wb.create_sheet(sheet)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c).value = _to_excel(value)
    wb.save(path)


# -----------------------------------------------------------------------------
# raw_data.xlsx helpers
# Layout of every sheet: Time Period | Year | then 5 groups of Max, Med, Min,
# 5 columns apart (C-E, H-J, M-O, R-T, W-Y). Groups 1-4 are the scenarios,
# group 5 is the government plan (where a model has one).
# -----------------------------------------------------------------------------
RAW_SCENARIOS = ['ConstCurrent', 'RagingStorm', 'RisingTide', 'ShiftingWinds']
RAW_COLUMN_OFFSET = 5
RAW_FIRST_YEAR = 2022.5  # Period 1 of the half-year grid used for hydrogen, PV and Wind
RAW_NUM_PERIODS = 36     # 2022.5 ... 2040


def raw_header():
    header = ['Time Period', 'Year'] + [None] * (RAW_COLUMN_OFFSET * 5)
    for s in range(5):
        for j, level in enumerate(['Max', 'Med', 'Min']):
            name = f'{RAW_SCENARIOS[s]} {level}' if s < 4 else f'{level}_Employment_S5'
            header[2 + s * RAW_COLUMN_OFFSET + j] = name
    return header


def raw_rows_from_output_table(output_rows):
    """A model output table (Time Period, Year, Max/Med/Min_Employment_S1...) with the raw_data headers."""
    header = raw_header()
    rows = [header]
    for row in output_rows[1:]:
        rows.append(list(row) + [None] * max(0, len(header) - len(row)))
    return rows


def raw_rows_by_year(years, groups):
    """Put results that have their own years onto the 36 half-year periods (2022.5-2040).

    groups: list of (max, med, min) arrays, one per scenario, aligned with years.
    Periods without results are 0; the Year column shows whole years only.
    """
    position = {round(float(y) * 2): i for i, y in enumerate(years)}
    rows = [raw_header()]
    for p in range(1, RAW_NUM_PERIODS + 1):
        year = RAW_FIRST_YEAR + 0.5 * (p - 1)
        row = [p, int(year) if year.is_integer() else None] + [None] * (RAW_COLUMN_OFFSET * 5)
        i = position.get(round(year * 2))
        for s, group in enumerate(groups):
            for j in range(3):
                row[2 + s * RAW_COLUMN_OFFSET + j] = 0 if i is None else group[j][i]
        rows.append(row)
    return rows


def write_raw_data_sheet(sheet, rows):
    write_cells(RAW_DATA_FILE, sheet, rows, replace_sheet=True)
    print(f'raw_data.xlsx: sheet "{sheet}" updated')


# -----------------------------------------------------------------------------
# Chart helper shared by the steel/aluminium/cement and manufacturing cells
# -----------------------------------------------------------------------------
def plot_scenario_bands(plot_years, y_max, y_med, y_min, plot_year_indices, names, colors, title,
                        integer_ticks=False, xticks=None):
    """Median line, shaded min-max band and dashed min/max lines per scenario (12 x 8 cm figure)."""
    fig, ax = plt.subplots(figsize=(12 / 2.54, 8 / 2.54))
    for s in range(len(y_max)):
        ax.fill_between(plot_years, y_min[s][plot_year_indices], y_max[s][plot_year_indices],
                        color=colors[s], alpha=0.2, edgecolor='none')
        ax.plot(plot_years, y_med[s][plot_year_indices], color=colors[s], linewidth=1.5, label=names[s])
        ax.plot(plot_years, y_max[s][plot_year_indices], '--', color=colors[s], linewidth=1)
        ax.plot(plot_years, y_min[s][plot_year_indices], '--', color=colors[s], linewidth=1)
    ax.set_title(title, fontsize=10, fontweight='normal')
    ax.set_xlabel('Years', fontsize=11)
    ax.set_ylabel('Number of Jobs', fontsize=12)
    if xticks is not None:
        ax.set_xticks(xticks)
    if integer_ticks:
        ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f'{v:.0f}'))
    ax.tick_params(labelsize=10)
    ax.legend(loc='upper center', bbox_to_anchor=(0.5, -0.2), ncol=2, frameon=False, fontsize=10)
    ax.grid(True)
    fig.tight_layout()


print('Setup complete.')
print(f'  Input workbook: {INPUT_FILE}')
print(f'  Results folder: {OUTPUT_DIR}')
